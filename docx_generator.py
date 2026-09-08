"""
DOCX + PDF export for LinkedIn Articles.

- build_docx(): DOCX book from article dicts (same shape as build_epub
  expects: title/url/published/content/profile/banner_image). This is the
  editable format for Google Docs: download the .docx and import it.
- docx_to_pdf(): convert the DOCX to PDF with headless LibreOffice.

Cover: optional image bytes (see cover_art.prepare_cover). The whole image
is used as the first page, nothing overlaid on it.
"""

import base64
import re
import shutil
import subprocess
import tempfile
from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Optional

from bs4 import BeautifulSoup, NavigableString
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

# Usable widths (A4 with 2 cm margins ≈ 17 cm text width).
BODY_IMG_WIDTH = Inches(6.0)
BANNER_WIDTH = Inches(6.0)

_CODE_BG = "1A3A52"  # dark blue, matches the ePub stylesheet


# Body spacing: exactly one blank line between paragraphs
# (space_after == one line pitch, measured in the PDF output).
_BODY_AFTER = Pt(15)


# ── Small python-docx helpers ────────────────────────────────────────────────

def _set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement("w:tcBorders")
        tcPr.append(tcBorders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tcBorders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tcBorders.append(element)
        for key, val in kwargs.items():
            element.set(qn(key), val)


def _shade_paragraph(paragraph, fill: str):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


# Inside padding for code blocks (paragraph shading has no padding control,
# so code goes in a borderless single cell whose margins pad all sides).
_CODE_CELL_MARGIN_DXA = 150  # ≈ 0.1"


def _shade_cell(cell, fill: str):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def _set_cell_margins(cell, dxa: int):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:w"), str(dxa))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tcPr.append(mar)


def _borderless(table):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "nil")
        el.set(qn("w:sz"), "0")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "auto")
        borders.append(el)
    tblPr.append(borders)


def _add_code_block(doc: Document, text: str):
    """Shaded code block with true inside padding + blank line after."""
    table = doc.add_table(rows=1, cols=1)
    _borderless(table)
    cell = table.cell(0, 0)
    _shade_cell(cell, _CODE_BG)
    _set_cell_margins(cell, _CODE_CELL_MARGIN_DXA)
    para = cell.paragraphs[0]
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lines = text.splitlines() or [""]
    for n, line in enumerate(lines):
        if n > 0:
            para.add_run().add_break()
        run = para.add_run(line if line else " ")
        _style_code_run(run)
    # Blank line after the block (tables take no paragraph spacing).
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.space_after = Pt(0)


def _add_hyperlink(paragraph, url: str, text: str):
    """Add a clickable hyperlink run to a paragraph."""
    part = paragraph.part
    r_id = part.relate_to(url, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0A66C2")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


_bookmark_counter = [0]


def _add_bookmark(paragraph, name: str):
    """Wrap a paragraph's runs in a bookmark so the TOC can link to it."""
    _bookmark_counter[0] += 1
    bid = str(_bookmark_counter[0])
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), bid)
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), bid)
    p = paragraph._p
    p.insert(0, start)
    p.append(end)


def _add_internal_hyperlink(paragraph, anchor: str, text: str):
    """Clickable link to a bookmark in the same document (TOC entries)."""
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("w:anchor"), anchor)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0A66C2")
    rPr.append(color)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.set(qn("xml:space"), "preserve")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _add_field(paragraph, instruction: str, font_size=Pt(9)):
    """Append a field run (PAGE / NUMPAGES) evaluated on open/export."""
    run = paragraph.add_run()
    run.font.size = font_size
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


def _add_page_number_footer(doc: Document):
    """Centered 'Page X of Y' footer; the cover page stays unnumbered."""
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    footer = section.footer
    para = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label = para.add_run("Page ")
    label.font.size = Pt(9)
    label.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    _add_field(para, "PAGE")
    mid = para.add_run(" of ")
    mid.font.size = Pt(9)
    mid.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    _add_field(para, "NUMPAGES")


def _style_code_run(run):
    run.font.name = "Consolas"
    run.font.size = Pt(9)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)


def _decode_data_uri(src: str) -> Optional[tuple[bytes, str]]:
    """Decode data:image/...;base64,... → (bytes, ext). None if not usable."""
    m = re.match(r"data:(image/\w+);base64,(.+)", (src or "").strip(), re.DOTALL)
    if not m:
        return None
    mime, b64 = m.group(1).lower(), m.group(2)
    ext = {"image/jpeg": "jpg", "image/png": "png",
           "image/gif": "gif", "image/webp": "webp"}.get(mime)
    if not ext:
        return None
    try:
        return base64.b64decode(b64), ext
    except Exception:
        return None


def _add_cover_image(doc: Document, cover_image: bytes):
    """Cover picture scaled to fit the whole page (never cropped)."""
    section = doc.sections[0]
    text_w = section.page_width - section.left_margin - section.right_margin
    text_h = section.page_height - section.top_margin - section.bottom_margin
    width = text_w
    try:
        with Image.open(BytesIO(cover_image)) as im:
            iw, ih = im.size
            if iw > 0 and ih > 0:
                width = min(text_w, int(text_h * iw / ih))
    except Exception:
        pass
    para = _add_image_paragraph(doc, cover_image, width)
    if para is not None:
        # No paragraph spacing around the cover: it must fit the page exactly.
        para.paragraph_format.space_before = Pt(0)
        para.paragraph_format.space_after = Pt(0)
    return para


def _add_image_paragraph(doc: Document, data: bytes, width):
    """Centered picture paragraph; returns it, or None when unusable."""
    try:
        Image.open(BytesIO(data)).verify()
    except Exception:
        return None
    try:
        para = doc.add_paragraph()
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        para.add_run().add_picture(BytesIO(data), width=width)
        return para
    except Exception:
        return None


# ── HTML → docx ──────────────────────────────────────────────────────────────

def _add_inline_runs(paragraph, element, link_url: Optional[str] = None):
    """Append runs for element's inline content (recurses into children)."""
    from bs4 import Comment as _CM
    for child in element.children:
        if isinstance(child, _CM):
            continue  # LinkedIn's empty <!-- --> separators: never a space run
        if isinstance(child, NavigableString):
            # Collapse whitespace: a raw "\n" would become a <w:br/>
            # (blank line) via add_run. Keep at most one separating space,
            # and never a leading one.
            text = re.sub(r"\s+", " ", str(child))
            if text.strip():
                paragraph.add_run(text)
            elif text == " " and paragraph.text:
                paragraph.add_run(" ")
        elif child.name in ("br",):
            paragraph.add_run().add_break()
        elif child.name in ("b", "strong", "i", "em", "u", "code"):
            # Recurse so inter-element spaces survive (stripped_strings
            # would eat them: "Escrevi<sp> </sp>" + link collapsed into
            # "Escreviesseartigo"), then style the added runs.
            before = len(paragraph.runs)
            _add_inline_runs(paragraph, child)
            for r2 in paragraph.runs[before:]:
                if child.name in ("b", "strong"):
                    r2.bold = True
                elif child.name in ("i", "em"):
                    r2.italic = True
                elif child.name == "u":
                    r2.underline = True
                else:
                    _style_code_run(r2)
        elif child.name == "a" and child.get("href"):
            label = child.get_text(" ", strip=True) or child["href"]
            _add_hyperlink(paragraph, child["href"], label)
        elif child.name == "img":
            pass  # images are hoisted to picture paragraphs at block level
        else:
            # span, small, font, unknown inline → recurse
            try:
                _add_inline_runs(paragraph, child)
            except Exception:
                paragraph.add_run(child.get_text(" ", strip=True))


def _add_quote_runs(paragraph, el):
    """Flatten a <blockquote> into one paragraph: inline runs + breaks."""
    from bs4 import Comment as _CM
    for child in el.children:
        if isinstance(child, _CM):
            continue
        if isinstance(child, NavigableString):
            text = " ".join(str(child).split())
            if text:
                paragraph.add_run(text + " ")
        elif getattr(child, "name", None) == "br":
            paragraph.add_run().add_break()
        elif getattr(child, "name", None) == "img":
            pass  # hoisted to a picture paragraph at block level
        elif getattr(child, "name", None) in ("ul", "ol"):
            for li in child.find_all("li", recursive=False):
                text = li.get_text(" ", strip=True)
                if text:
                    paragraph.add_run(text + " ")
            paragraph.add_run().add_break()
        else:
            try:
                text = child.get_text(" ", strip=True)
            except Exception:
                text = ""
            if text:
                paragraph.add_run(text + " ")


# Tags that must stand as their own blocks, never inline inside a paragraph.
_LIFT_BLOCKS = ["ul", "ol", "pre", "table", "blockquote",
                "h1", "h2", "h3", "figure", "hr", "div"]


def _lift_blocks(soup):
    """Repair sloppy HTML where blocks nest inside <p>/<li>.

    Browsers auto-close <p> before <ul> etc.; bs4 keeps the nesting, which
    would collapse a whole list into one bullet-less paragraph. So:
    - <p> containing blocks is unwrapped (children become siblings);
    - blocks nested directly in <li> are moved to after the <li>, so the
      item keeps its own bullet.
    """
    for p in soup.find_all("p"):
        if p.find(_LIFT_BLOCKS):
            p.unwrap()
    for li in soup.find_all("li"):
        anchor = li
        for blk in list(li.find_all(_LIFT_BLOCKS, recursive=False)):
            blk.extract()
            anchor.insert_after(blk)
            anchor = blk


def _add_hoisted_images(doc: Document, el):
    """Add picture paragraphs for data-URI <img> nested inside a block."""
    try:
        imgs = el.find_all("img")
    except Exception:
        return
    for img in imgs:
        decoded = _decode_data_uri(img.get("src", ""))
        if decoded:
            _add_image_paragraph(doc, decoded[0], BODY_IMG_WIDTH)


# Punctuation that must hug adjacent words (no space before closers,
# none after openers) when a whitespace-only node sits at a tag boundary.
_CLOSING_PUNCT = ",.;:!?%)]}'\"\"''"
_OPENING_PUNCT = "([{‘“\""
# Straight double-quote excluded: it opens as often as it closes, so it is
# resolved contextually inside _fix_punct_spacing instead.
_CLOSING_NO_DQ = _CLOSING_PUNCT.replace('"', "")
_OPENING_NO_DQ = _OPENING_PUNCT.replace('"', "")


_LINKEDIN_BASE = "https://www.linkedin.com"


def _strip_linkedin_comments(soup):
    """Remove LinkedIn's empty <!-- --> separators between inline nodes.

    A Comment stringifies to a single space, which _add_inline_runs would
    otherwise emit as a space run — the 'Name .' / 'bold .' gap.

    Walks .descendants manually: find_all(string=...) never descends into
    <pre> in BeautifulSoup, so code-block comments would survive it.
    """
    from bs4 import Comment as _CM
    for el in list(soup.descendants):
        if isinstance(el, _CM):
            el.extract()


def _absolutize_links(soup):
    """Rewrite relative LinkedIn hrefs (/in/..., ../in/..., ../../in/...)
    as absolute https://www.linkedin.com/... URLs so hyperlinks in the
    book resolve correctly outside linkedin.com."""
    from urllib.parse import urljoin
    for a in soup.find_all("a", href=True):
        href = (a.get("href") or "").strip()
        if not href:
            continue
        low = href.lower()
        if low.startswith(("#", "mailto:", "tel:", "data:", "javascript:")):
            continue
        if low.startswith(("http://", "https://")):
            continue
        if href.startswith("//"):
            a["href"] = "https:" + href
            continue
        a["href"] = urljoin(_LINKEDIN_BASE + "/", href)


def _tighten_punct_text(soup):
    """Collapse spaces before closing punctuation inside text nodes
    (e.g. 'operacional .' → 'operacional.'). Skips pre/code (verbatim)
    and whitespace-only nodes.

    The straight double-quote is ambiguous, so it is handled contextually:
    a closing one hugs (drop the space), an opening one keeps — or gains —
    exactly one space before it ('dei"match"' → 'dei "match"').
    """
    from bs4 import Comment as _CM
    _CLOSE_NO_DQ = _CLOSING_PUNCT.replace('"', "")
    _OPEN_NO_DQ = _OPENING_PUNCT.replace('"', "")
    for s in soup.find_all(string=True):
        if isinstance(s, _CM):
            continue
        if s.parent is not None and s.parent.name in ("pre", "code"):
            continue
        t = str(s)
        if not t or not t.strip():
            continue
        new = re.sub(r"\s+([%s])" % re.escape(_CLOSE_NO_DQ), r"\1", t)
        new = re.sub(r"([%s])\s+" % re.escape(_OPEN_NO_DQ), r"\1", new)
        # Closing straight quote: drop a preceding space.
        new = re.sub(r'\s+"(?=[\s.,;:!?%)\]}’”]|$)', '"', new)
        # Opening straight quote glued to a word: insert the missing space,
        # but only when a later quote closes the pair ('dei"match" x' →
        # 'dei "match" x'; an unmatched '"fim"disse' is left alone).
        qs = [m.start() for m in re.finditer(r'"', new)]
        out, last = [], 0
        for m in re.finditer(r'([^\s\(\[{‘“"\'’”])"(?=\w)', new):
            if any(q > m.end() - 1 for q in qs):
                out.append(new[last:m.start()])
                out.append(m.group(1) + ' "')
                last = m.end()
        out.append(new[last:])
        new = "".join(out)
        if new != t:
            s.replace_with(new)


def _fix_punct_spacing(soup):
    """Drop whitespace-only nodes glued to punctuation across tag boundaries.

    LinkedIn emits spaces as separate spans, sometimes before `,`/`.`/`)`
    or after `(`/quotes. Skips pre/code (verbatim) and only touches
    whitespace-only nodes, so emoticons and in-word spacing are safe.

    The straight double-quote is resolved contextually: the space goes
    away before a closing quote ('...Jobs ." ' → hug) and after an opening
    one, but is kept before an opening quote and after a closing one.
    """
    from bs4 import NavigableString as _NS, Comment as _CM

    def _meaningful(t):
        return (isinstance(t, _NS) and not isinstance(t, _CM)
                and bool(str(t).strip()))

    def _after_dquote(nxt):
        """Char following a leading '"' (peeks past lone quotes)."""
        rest = str(nxt).lstrip()[1:].lstrip()
        if rest:
            return rest[:1]
        fol = nxt.find_next(string=_meaningful)
        return str(fol).lstrip()[:1] if fol is not None else ""

    def _before_dquote(prev):
        """Char preceding a trailing '"' (peeks past lone quotes)."""
        rest = str(prev).rstrip()[:-1].rstrip()
        if rest:
            return rest[-1:]
        prv = prev.find_previous(string=_meaningful)
        return str(prv).rstrip()[-1:] if prv is not None else ""

    for s in soup.find_all(string=True):
        if isinstance(s, _CM):
            continue
        if s.parent is not None and s.parent.name in ("pre", "code"):
            continue
        if not s or str(s).strip():
            continue
        prev = s.find_previous(string=_meaningful)
        nxt = s.find_next(string=_meaningful)
        prev_ch = str(prev).rstrip()[-1:] if prev is not None else ""
        next_ch = str(nxt).lstrip()[:1] if nxt is not None else ""
        if next_ch in _CLOSING_NO_DQ or prev_ch in _OPENING_NO_DQ:
            drop = True
        elif next_ch == '"':
            after = _after_dquote(nxt)
            drop = after == "" or after in _CLOSING_NO_DQ
        elif prev_ch == '"':
            before = _before_dquote(prev)
            drop = before == "" or before in _OPENING_NO_DQ
        else:
            drop = False
        if drop:
            s.extract()


def _strip_edge_breaks(el):
    """Remove leading/trailing <br> children (phantom blank lines)."""
    try:
        children = list(el.children)
    except Exception:
        return
    while children and getattr(children[0], "name", None) == "br":
        children[0].extract()
        children = children[1:]
    while children and getattr(children[-1], "name", None) == "br":
        children[-1].extract()
        children = children[:-1]


def _add_block(doc: Document, el, list_style: Optional[str] = None):
    """Convert one block-level bs4 element into docx content."""
    from bs4 import Comment as _CM
    if isinstance(el, _CM):
        return  # never render comment separators as text
    name = getattr(el, "name", None)
    if name is None:  # bare string
        text = str(el).strip()
        if text:
            doc.add_paragraph(text, style=list_style or "Normal")
        return

    if name in ("h1", "h2", "h3"):
        level = {"h1": 1, "h2": 2, "h3": 3}[name]
        _strip_edge_breaks(el)
        doc.add_heading(el.get_text(" ", strip=True), level=level)
    elif name == "p":
        if not el.get_text(strip=True) and not el.find("img"):
            return  # skip blank paragraphs: no double gaps
        _strip_edge_breaks(el)
        para = doc.add_paragraph(style=list_style or "Normal")
        _add_inline_runs(para, el)
        _add_hoisted_images(doc, el)
    elif name == "pre":
        text = el.get_text("\n")
        _add_code_block(doc, text)
    elif name in ("ul", "ol"):
        style = "List Bullet" if name == "ul" else "List Number"
        for li in el.find_all("li", recursive=False):
            _add_block(doc, li, list_style=style)
            for sub in li.find_all(["ul", "ol"], recursive=False):
                _add_block(doc, sub)
    elif name == "li":
        _strip_edge_breaks(el)
        para = doc.add_paragraph(style=list_style or "List Bullet")
        _add_inline_runs(para, el)
        _add_hoisted_images(doc, el)
    elif name == "blockquote":
        # One single paragraph for the whole quote (mixed italic/normal
        # text must not gain extra line spacing). "Quote" style: same
        # spacing as body but left-aligned, never justified.
        para = doc.add_paragraph(style="BookQuote")
        _add_quote_runs(para, el)
        for run in para.runs:
            run.italic = True
        _add_hoisted_images(doc, el)
    elif name == "img":
        decoded = _decode_data_uri(el.get("src", ""))
        if decoded:
            _add_image_paragraph(doc, decoded[0], BODY_IMG_WIDTH)
    elif name == "figure":
        for child in el.children:
            _add_block(doc, child)
    elif name == "figcaption":
        para = doc.add_paragraph(style="BookCaption")
        para.add_run(el.get_text(" ", strip=True))
    elif name == "table":
        rows = el.find_all("tr")
        if not rows:
            return
        table = doc.add_table(rows=0, cols=0)
        table.style = "Table Grid"
        for tr in rows:
            cells = tr.find_all(["th", "td"])
            if not cells:
                continue
            row = table.add_row()
            # grow columns as needed
            while len(row.cells) < len(cells):
                table.add_column()
            for cell_el, cell in zip(cells, row.cells):
                cell.text = ""
                para = cell.paragraphs[0]
                para.paragraph_format.space_after = Pt(0)  # compact cells
                para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
                _add_inline_runs(para, cell_el)
                if cell_el.name == "th":
                    for run in para.runs:
                        run.bold = True
    elif name == "hr":
        doc.add_paragraph("─" * 40, style="Normal")
    elif name in ("div", "section", "article", "header", "footer", "main"):
        for child in el.children:
            _add_block(doc, child)
    else:
        # Unknown block-ish tag: inline content as a paragraph.
        text = el.get_text(" ", strip=True) if hasattr(el, "get_text") else ""
        if text:
            para = doc.add_paragraph(style=list_style or "Normal")
            _add_inline_runs(para, el)


def _add_body_html(doc: Document, content: str):
    """Append stored article HTML (or plain text) to the document."""
    if not content or not content.strip():
        doc.add_paragraph("No content available.", style="Normal").runs[0].italic = True
        return
    if content.strip().startswith("<"):
        soup = BeautifulSoup(content, "html.parser")
        _strip_linkedin_comments(soup)
        _absolutize_links(soup)
        _lift_blocks(soup)
        _fix_punct_spacing(soup)
        _tighten_punct_text(soup)
        try:
            from scraper import _separate_glued_quotes as _sep_q
        except ImportError:  # pragma: no cover
            _sep_q = None
        if _sep_q is not None:
            _sep_q(soup)
        root = soup.body or soup
        from bs4 import Comment as _CM
        for child in list(root.children):
            if isinstance(child, _CM):
                continue
            try:
                _add_block(doc, child)
            except Exception:
                continue
        return
    for para in re.split(r"\n{2,}", content.strip()):
        text = " ".join(l.strip() for l in para.splitlines() if l.strip())
        if text:
            doc.add_paragraph(text, style="Normal")


def _banner_bytes(art: dict) -> Optional[bytes]:
    """Decode the stored base64 banner_image, if present."""
    b64 = art.get("banner_image") or ""
    if not b64.startswith("data:"):
        return None
    decoded = _decode_data_uri(b64)
    return decoded[0] if decoded else None


# ── Final BrodTec page ─────────────────────────────────────────────────────────

_BRODTEC_URL = "https://brodtec.com"


def _load_brodtec_logo() -> Optional[bytes]:
    """Load the BrodTec logo for offline embedding. None when missing."""
    from pathlib import Path as _Path
    try:
        data = (_Path(__file__).parent / "web" / "static" / "logobrodlink.png").read_bytes()
        return bytes(data) or None
    except Exception:
        return None


def _add_brodtec_page(doc: Document):
    """Append the final one-page pt-BR summary of brodtec.com.

    Mirrors the landing-page layout top-down; the contact form is replaced
    by plain contact links (a form makes no sense on paper).
    """
    doc.add_page_break()
    logo = _load_brodtec_logo()
    if logo:
        _add_image_paragraph(doc, logo, Inches(2.0))

    doc.add_heading("Sobre a BrodTec", level=1)
    hero = doc.add_paragraph(style="Normal")
    hero.add_run("Transforme sua tecnologia de um centro de custo em um motor de receita.").bold = True
    doc.add_paragraph(
        "Ajudamos organizações a desescalar a complexidade, cruzar fronteiras e "
        "libertar talentos humanos por meio do uso estratégico da Inteligência Artificial.",
        style="Normal",
    )

    doc.add_heading("O problema", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("A tecnologia virou uma caixa-preta.").bold = True
    p.add_run(" Os investimentos entram, mas os resultados demoram, os prazos estouram "
              "e os custos com licenças e fornecedores só aumentam.")
    for q in ("Sua TI parece um freio de mão em vez de um acelerador?",
              "Sua equipe está sobrecarregada com tarefas repetitivas que não geram valor?",
              "Sua expansão para novos mercados está travada por barreiras culturais e técnicas?"):
        doc.add_paragraph(q, style="List Bullet")
    doc.add_paragraph(
        "Nós resolvemos a cegueira de eficiência: tiramos a tecnologia do isolamento "
        "operacional e a reposicionamos como parceira estratégica do faturamento.",
        style="Normal",
    )

    doc.add_heading("Para quem", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Agilidade, não burocracia.").bold = True
    for title, desc in (
        ("Pequenas e Médias Empresas",
         "automatizar processos manuais e adotar IA para crescer sem explodir os custos."),
        ("Empresas de Médio Porte",
         "deixar de ser refém de sistemas engessados e otimizar o que já existe."),
        ("Organizações em Expansão",
         "entrar no mercado brasileiro ou levar sua tecnologia ao exterior sem perder a coesão."),
    ):
        pp = doc.add_paragraph(style="Normal")
        pp.add_run(title + " — ").bold = True
        pp.add_run(desc)

    doc.add_heading("Como funciona", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Baseado no que sua empresa já tem.").bold = True
    p.add_run(" Sem fórmulas prontas.")
    for title, desc in (
        ("01 — Diagnóstico Rápido (Blitz).",
         "Em poucas semanas, identificamos gargalos invisíveis e entregamos um protótipo "
         "funcional de IA personalizado para o seu negócio."),
        ("02 — Setup de Eficiência.",
         "Intervenção direta para renegociar contratos, melhorar fluxos de entrega "
         "e reduzir custos fixos."),
        ("03 — Arquitetura de Pontes.",
         "Orquestramos a transição cultural e técnica para operar com agilidade "
         "em qualquer território."),
    ):
        pp = doc.add_paragraph(style="Normal")
        pp.add_run(title + " ").bold = True
        pp.add_run(desc)

    doc.add_heading("Resultados", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Transformações tangíveis.").bold = True
    p.add_run(" O custo de não agir costuma ser muito maior que o investimento.")
    for title, desc in (
        ("Economia Real", "redução imediata de custos operacionais."),
        ("Velocidade de Entrega", "o dobro de entregas com a mesma equipe."),
        ("Agência Humana", "pessoas focadas na estratégia; o trabalho braçal fica com a IA."),
        ("Previsibilidade", "parar de \u201cestimar\u201d e começar a \u201corçar\u201d com base na realidade."),
    ):
        pp = doc.add_paragraph(style="List Bullet")
        pp.add_run(title + " — ").bold = True
        pp.add_run(desc)

    doc.add_heading("Por que confiar", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Mais de 30 anos domesticando a complexidade.").bold = True
    p.add_run(" Cesar Brod, fundador da BrodTec, liderou a entrada de gigantes como Tandem "
              "e ACI Worldwide no Brasil, expandiu operações para mais de 10 países e criou "
              "a primeira cooperativa de software livre do mundo. Autor e tradutor de obras "
              "de referência em engenharia de software.")

    doc.add_heading("Como começar", level=2)
    p = doc.add_paragraph(style="Normal")
    p.add_run("Workshop de Liderança e IA").bold = True
    p.add_run(" — intervenção prática de quatro horas: diagnóstico inicial de eficiência e a "
              "primeira versão de um assistente de IA personalizado. "
              "Investimento inicial: R$ 2.000,00.")

    contact = doc.add_paragraph(style="Normal")
    contact.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_hyperlink(contact, "https://brodtec.com", "brodtec.com")
    contact.add_run("  ·  ")
    _add_hyperlink(contact, "https://wa.me/5551981361214", "WhatsApp +55 51 98136-1214")
    contact.add_run("  ·  ")
    _add_hyperlink(contact, "mailto:cesar@brodtec.com", "cesar@brodtec.com")
    contact.add_run("  ·  ")
    _add_hyperlink(contact, "https://www.linkedin.com/in/cesarbrod/", "LinkedIn")

    footer = doc.add_paragraph(style="Normal")
    footer.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("© 2026 BrodTec")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


# ── Main builder ─────────────────────────────────────────────────────────────

def build_docx(
    articles: list,           # same shape as build_epub expects
    title: str = "LinkedIn Articles",
    author: str = "LinkedIn Articles Export",
    cover_image: Optional[bytes] = None,   # validated via cover_art
    output_path: Optional[Path] = None,
    brodtec_page: bool = True,   # append final "Sobre a BrodTec" page
) -> Path:
    """Build a .docx book. Returns the output path."""
    if output_path is None:
        from datetime import datetime as _dt
        ts = _dt.now().strftime("%Y%m%d_%H%M%S")
        output_path = Path(__file__).parent / f"export_{ts}.docx"
    output_path = Path(output_path)

    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.27)    # A4
        section.page_height = Inches(11.69)
        section.top_margin = section.bottom_margin = Inches(0.8)
        section.left_margin = section.right_margin = Inches(0.8)
    _add_page_number_footer(doc)

    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    # Regular paragraphs are justified; everything else gets an explicit
    # alignment below so headings, lists, quotes and captions stay put.
    style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    # Single spacing between paragraphs: no extra gap before/after,
    # exactly one line break worth of separation.
    style.paragraph_format.space_before = Pt(0)
    style.paragraph_format.space_after = _BODY_AFTER
    style.paragraph_format.line_spacing = 1.0
    for list_style_name in ("List Bullet", "List Number"):
        try:
            lst = doc.styles[list_style_name]
            lst.paragraph_format.space_before = Pt(0)
            lst.paragraph_format.space_after = _BODY_AFTER
            lst.paragraph_format.line_spacing = 1.0
        except KeyError:
            pass
    # Headings keep a hint of air (hierarchy), but not the huge defaults.
    for heading_style_name in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            hst = doc.styles[heading_style_name]
            hst.paragraph_format.space_before = _BODY_AFTER
            hst.paragraph_format.space_after = Pt(8)
            hst.paragraph_format.line_spacing = 1.0
            hst.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        except KeyError:
            pass
    for list_style_name in ("List Bullet", "List Number"):
        try:
            doc.styles[list_style_name].paragraph_format.alignment = (
                WD_ALIGN_PARAGRAPH.LEFT)
        except KeyError:
            pass
    try:
        doc.styles["Title"].paragraph_format.alignment = (
            WD_ALIGN_PARAGRAPH.CENTER)
    except KeyError:
        pass

    # Dedicated styles so quotes/captions are not affected by justification.
    from docx.enum.style import WD_STYLE_TYPE
    quote_style = doc.styles.add_style("BookQuote", WD_STYLE_TYPE.PARAGRAPH)
    quote_style.base_style = doc.styles["Normal"]
    quote_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    quote_style.paragraph_format.left_indent = Inches(0.4)
    caption_style = doc.styles.add_style("BookCaption", WD_STYLE_TYPE.PARAGRAPH)
    caption_style.base_style = doc.styles["Normal"]
    caption_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_style.font.italic = True
    caption_style.font.size = Pt(9)

    # ── Cover: whole image, nothing overlaid ──
    if cover_image:
        if _add_cover_image(doc, cover_image):
            doc.add_page_break()

    # ── Title page ──
    doc.add_heading(title, level=0)
    byline = doc.add_paragraph(f"By {author}", style="Normal")
    byline.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    dateline = doc.add_paragraph(
        f"Generated {datetime.now().strftime('%B %d, %Y')} — "
        f"{len(articles)} article(s)",
        style="Normal",
    )
    dateline.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_page_break()

    # ── Table of contents (manual entries with internal links) ──
    # A real TOC field stays empty in headless LibreOffice PDF exports,
    # so entries link directly to chapter bookmarks instead. Page numbers
    # are omitted: pagination is only known when the file is opened.
    if articles:
        doc.add_heading("Contents", level=1)
        for idx, row in enumerate(articles, 1):
            art = dict(row) if hasattr(row, "keys") else row
            entry = doc.add_paragraph(style="Normal")
            entry.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
            _add_internal_hyperlink(
                entry, f"chap_{idx:04d}", art.get("title") or f"Article {idx}")
        doc.add_page_break()

    for idx, row in enumerate(articles, 1):
        art = dict(row) if hasattr(row, "keys") else row
        art_title = art.get("title") or f"Article {idx}"
        art_url = art.get("url") or ""
        art_published = art.get("published") or "unknown date"
        art_profile = art.get("profile") or ""
        art_content = art.get("content") or ""

        if idx > 1:
            doc.add_page_break()
        heading = doc.add_heading(art_title, level=1)
        _add_bookmark(heading, f"chap_{idx:04d}")

        meta = doc.add_paragraph(style="Normal")
        meta.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        run = meta.add_run(f"By {art_profile}  •  {art_published}  •  ")
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        if art_url:
            _add_hyperlink(meta, art_url, "View on LinkedIn ↗")

        banner = _banner_bytes(art)
        if banner:
            _add_image_paragraph(doc, banner, BANNER_WIDTH)

        _add_body_html(doc, art_content)

    if brodtec_page:
        _add_brodtec_page(doc)

    doc.save(str(output_path))
    return output_path


# ── DOCX → PDF via headless LibreOffice ──────────────────────────────────────

def docx_to_pdf(docx_path: Path, out_dir: Optional[Path] = None,
                timeout: int = 180) -> Path:
    """Convert a .docx to .pdf with `soffice --headless`. Returns pdf path."""
    docx_path = Path(docx_path)
    if shutil.which("soffice") is None and shutil.which("libreoffice") is None:
        raise RuntimeError("LibreOffice (soffice) is not installed.")
    out_dir = Path(out_dir) if out_dir else docx_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    cmd = ["soffice", "--headless", "--convert-to", "pdf",
           "--outdir", str(out_dir), str(docx_path)]
    try:
        proc = subprocess.run(cmd, capture_output=True, text=True,
                              timeout=timeout)
    except subprocess.TimeoutExpired:
        raise RuntimeError("LibreOffice conversion timed out.")
    pdf_path = out_dir / (docx_path.stem + ".pdf")
    if proc.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(
            f"PDF conversion failed: {(proc.stderr or proc.stdout).strip()[-500:]}")
    return pdf_path
